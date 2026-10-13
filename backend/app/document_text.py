"""Text extraction from persisted bytes; client extracted_text is never release evidence."""
from __future__ import annotations

import hashlib
import io
import json
import resource
import subprocess
import sys
import zipfile
from defusedxml import ElementTree
from dataclasses import dataclass
from importlib.metadata import version
from pathlib import Path


MAX_FILE_BYTES = 16 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
MAX_PAGES = 200


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    document_sha256: str
    method: str

    @property
    def evidence(self) -> dict[str, str]:
        return {"document_sha256": self.document_sha256, "text_sha256": hashlib.sha256(self.text.encode()).hexdigest(), "method": self.method}


def extract_bytes(payload: bytes, suffix: str) -> tuple[str, str]:
    if not payload or len(payload) > MAX_FILE_BYTES:
        raise ValueError("Document size is outside extraction limits")
    if suffix in {".txt", ".md"}:
        if payload.startswith((b"%PDF-", b"PK\x03\x04")):
            raise ValueError("Stored format and filename disagree")
        text, method = payload.decode("utf-8"), "utf8-v1"
    elif suffix == ".pdf":
        from pypdf import PdfReader

        if not payload.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF header")
        reader = PdfReader(io.BytesIO(payload), strict=True)
        if reader.is_encrypted or not 0 < len(reader.pages) <= MAX_PAGES:
            raise ValueError("Encrypted, empty or oversized PDF is unsupported")
        if reader.trailer["/Root"].get("/AcroForm"):
            raise ValueError("PDF forms require a separate extraction contract")
        parts = []
        for page in reader.pages:
            resources = page.get("/Resources")
            resources = resources.get_object() if resources else {}
            if resources.get("/XObject") or resources.get("/ExtGState") or page.get("/Annots"):
                raise ValueError("PDF images, embedded forms and annotations require review")
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("PDF content stream exceeds extraction limits")
            allowed = {b"BT", b"ET", b"Tf", b"Td", b"TD", b"Tm", b"T*", b"Tj", b"TJ", b"'", b'"', b"Tc", b"Tw", b"Tz", b"TL", b"Ts", b"q", b"Q", b"cm"}
            for operands, operator in contents.operations if contents is not None else []:
                if operator == b"Tr" and operands == [0]:
                    continue
                if operator not in allowed:
                    raise ValueError("PDF rendering, clipping, color or optional content is unsupported")
            part = page.extract_text()
            if not part or not part.strip():
                raise ValueError("Every PDF page must contain extractable text")
            parts.append(part.rstrip())
        text, method = "\n".join(parts), f"pypdf-{version('pypdf')}-text-v1"
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX expands beyond extraction limits")
            forbidden = ("word/media/", "word/embeddings/", "word/header", "word/footer", "word/footnotes", "word/endnotes", "word/comments", "word/vbaProject")
            if any(e.filename.startswith(forbidden) or e.flag_bits & 1 for e in entries):
                raise ValueError("DOCX images, headers, notes, macros and encryption require review")
            unsupported = {"altChunk", "ins", "del", "txbxContent", "object", "vanish", "drawing", "pict", "sdt", "fldSimple", "fldChar"}
            for entry in entries:
                if entry.filename.endswith((".xml", ".rels")):
                    tree = ElementTree.fromstring(archive.read(entry), forbid_dtd=True, forbid_entities=True, forbid_external=True)
                    if any(element.tag.rsplit("}", 1)[-1] in unsupported for element in tree.iter()):
                        raise ValueError("DOCX external content, revisions and hidden styles are unsupported")
                    if any(element.attrib.get("TargetMode") == "External" for element in tree.iter()):
                        raise ValueError("DOCX external relationships are unsupported")
        document = Document(io.BytesIO(payload))
        parts = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            elif isinstance(block, Table):
                for row in block.rows:
                    if any(cell.tables for cell in row.cells):
                        raise ValueError("Nested DOCX tables are unsupported")
                    parts.append("\t".join(cell.text for cell in row.cells))
        text, method = "\n".join(parts), f"python-docx-{version('python-docx')}-body-v1"
    else:
        raise ValueError("Unsupported stored document format")
    if not text.strip() or len(text) > MAX_TEXT_CHARS or "\x00" in text:
        raise ValueError("Extracted text is empty or outside extraction limits")
    return text, method


def extract_document(path: Path, file_name: str, expected_sha256: str) -> ExtractedDocument:
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Document exceeds extraction limits")
    with path.open("rb") as stream:
        payload = stream.read(MAX_FILE_BYTES + 1)
    if len(payload) > MAX_FILE_BYTES:
        raise ValueError("Document exceeds extraction limits")
    actual = hashlib.sha256(payload).hexdigest()
    if actual != expected_sha256:
        raise ValueError("Stored document content changed")
    suffix = Path(file_name).suffix.lower()
    if suffix in {".txt", ".md"}:
        text, method = extract_bytes(payload, suffix)
    else:
        # The worker receives exactly the bytes hashed above, not a mutable path.
        result = subprocess.run(
            [sys.executable, "-m", "backend.app.document_text", suffix], input=payload,
            capture_output=True, timeout=12, check=False,
        )
        if result.returncode:
            raise ValueError("Stored document could not be safely extracted")
        decoded = json.loads(result.stdout)
        text, method = decoded["text"], decoded["method"]
    return ExtractedDocument(text=text, document_sha256=actual, method=method)


if __name__ == "__main__":
    resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
    try:
        data = sys.stdin.buffer.read(MAX_FILE_BYTES + 1)
        text, method = extract_bytes(data, sys.argv[1])
        print(json.dumps({"text": text, "method": method}, ensure_ascii=False))
    except Exception:
        sys.exit(1)
