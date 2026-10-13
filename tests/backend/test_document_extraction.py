from __future__ import annotations

import base64
import hashlib
import importlib
import io
from types import SimpleNamespace

import pytest
from fastapi import HTTPException


def pdf_bytes(text: str):
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    font = writer._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (" + text.encode("ascii") + b") Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_bytes(text: str):
    from docx import Document
    document = Document()
    document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.mark.parametrize("suffix,create", [(".pdf", pdf_bytes), (".docx", docx_bytes)])
def test_server_extracts_binary_formats_and_binds_original_bytes(app_bundle, make_user, suffix, create):
    schemas = importlib.import_module("backend.app.schemas.expedientes")
    service = importlib.import_module("backend.app.release_service")
    text = "Synthetic document with no legal claims."
    payload = create(text)
    actor = make_user(email="format@example.com", display_name="Reviewer", role=app_bundle["enums"].UserRole.ADMIN)["user"]
    request = SimpleNamespace(url=SimpleNamespace(path="/test"), method="POST")
    with app_bundle["db"].SessionLocal() as db:
        expediente = app_bundle["expediente_routes"].create_expediente(payload=schemas.ExpedienteCreate(title="Extraction", matter_type="amparo"), request=request, current_user=actor, db=db)
        detail = app_bundle["expediente_routes"].create_document(
            expediente_id=str(expediente.id), request=request, current_user=actor, db=db,
            payload=schemas.DocumentWithVersionCreate(document={"document_type": "escrito", "title": "Document"}, version={"source_kind": "uploaded", "file_name": "document" + suffix, "content_base64": base64.b64encode(payload).decode(), "extracted_text": "Falsified client extraction"}),
        )
        version = detail.documents[0].versions[0]
        version_id, storage_key, expediente_id = version.id, version.storage_key, expediente.id
        forged_bundle = {"response_type": "recommendation", "response_text": "Falsified client extraction", "citations": [], "claims": []}
        with pytest.raises(HTTPException):
            service.evaluate_release(db=db, actor_id=actor.id, expediente_id=expediente.id, document_version_id=version.id, bundle=forged_bundle)
        bundle = {**forged_bundle, "response_text": text}
        execution = service.evaluate_release(db=db, actor_id=actor.id, expediente_id=expediente.id, document_version_id=version.id, bundle=bundle)
        execution_id = execution.id
        assert execution.document_sha256 == hashlib.sha256(payload).hexdigest()
        assert execution.evidence["extraction"]["text_sha256"] == hashlib.sha256(text.encode()).hexdigest()
        receipt = service.release_from_execution(db=db, actor_id=actor.id, expediente_id=expediente.id, execution_id=execution.id, bundle=bundle)
        assert receipt.verdict.value == "approved"
    storage = importlib.import_module("backend.app.storage")
    storage.resolve_storage_path(storage_key).write_bytes(create("Different content"))
    with app_bundle["db"].SessionLocal() as db:
        retried = service.release_from_execution(db=db, actor_id=actor.id, expediente_id=expediente_id, execution_id=execution_id, bundle=bundle)
        assert retried.id == receipt.id
        with pytest.raises(HTTPException):
            service.evaluate_release(db=db, actor_id=actor.id, expediente_id=expediente_id, document_version_id=version_id, bundle=bundle)


@pytest.mark.parametrize("suffix,payload", [(".pdf", b"broken PDF"), (".docx", b"broken ZIP"), (".txt", b"%PDF-1.4 fake")])
def test_malformed_or_mislabeled_formats_fail_closed(app_bundle, tmp_path, suffix, payload):
    extraction = importlib.import_module("backend.app.document_text")
    path = tmp_path / ("broken" + suffix)
    path.write_bytes(payload)
    with pytest.raises(ValueError):
        extraction.extract_document(path, path.name, hashlib.sha256(payload).hexdigest())


def test_blank_and_encrypted_pdf_require_review(app_bundle, tmp_path):
    from pypdf import PdfWriter
    extraction = importlib.import_module("backend.app.document_text")
    for encrypted in (False, True):
        writer = PdfWriter()
        writer.add_blank_page(600, 800)
        if encrypted:
            writer.encrypt("synthetic-test-only")
        out = io.BytesIO()
        writer.write(out)
        payload = out.getvalue()
        path = tmp_path / "document.pdf"
        path.write_bytes(payload)
        with pytest.raises(ValueError):
            extraction.extract_document(path, path.name, hashlib.sha256(payload).hexdigest())


def test_docx_unsupported_stories_are_rejected(app_bundle, tmp_path):
    from docx import Document
    extraction = importlib.import_module("backend.app.document_text")
    document = Document()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Unreviewed header"
    out = io.BytesIO()
    document.save(out)
    payload = out.getvalue()
    path = tmp_path / "document.docx"
    path.write_bytes(payload)
    with pytest.raises(ValueError):
        extraction.extract_document(path, path.name, hashlib.sha256(payload).hexdigest())


def test_hidden_docx_style_is_rejected(app_bundle, tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    document = Document()
    style = document.styles.add_style("HiddenSyntheticStyle", WD_STYLE_TYPE.CHARACTER)
    style.font.hidden = True
    document.add_paragraph().add_run("Hidden text").style = style
    out = io.BytesIO()
    document.save(out)
    payload = out.getvalue()
    path = tmp_path / "hidden.docx"
    path.write_bytes(payload)
    with pytest.raises(ValueError):
        importlib.import_module("backend.app.document_text").extract_document(path, path.name, hashlib.sha256(payload).hexdigest())


@pytest.mark.parametrize("extra", [b"3 Tr", b"1 1 1 rg", b"/OC /Hidden BDC", b"/Invisible gs"])
def test_pdf_unsupported_rendering_is_rejected(app_bundle, tmp_path, extra):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject
    writer = PdfWriter()
    page = writer.add_page(PdfReader(io.BytesIO(pdf_bytes("Hidden text"))).pages[0])
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf " + extra + b" 72 720 Td (Hidden text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(contents)
    out = io.BytesIO()
    writer.write(out)
    payload = out.getvalue()
    path = tmp_path / "hidden.pdf"
    path.write_bytes(payload)
    with pytest.raises(ValueError):
        importlib.import_module("backend.app.document_text").extract_document(path, path.name, hashlib.sha256(payload).hexdigest())


def test_pdf_indirect_resources_are_supported(app_bundle, tmp_path):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import NameObject
    writer = PdfWriter()
    page = writer.add_page(PdfReader(io.BytesIO(pdf_bytes("Synthetic text"))).pages[0])
    page[NameObject("/Resources")] = writer._add_object(page["/Resources"])
    out = io.BytesIO()
    writer.write(out)
    payload = out.getvalue()
    path = tmp_path / "simple.pdf"
    path.write_bytes(payload)
    result = importlib.import_module("backend.app.document_text").extract_document(path, path.name, hashlib.sha256(payload).hexdigest())
    assert result.text == "Synthetic text"
